from docx import Document
from docx.oxml.ns import qn

doc = Document(r'C:\Users\qiucanro\pygodAppData\content_files\bulletin_2026-3.docx')

print("=== BOLD TEXT IN PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    bold_runs = [(r.text, r.font.bold) for r in p.runs if r.font.bold]
    normal_runs = [(r.text, r.font.bold) for r in p.runs if not r.font.bold]
    if bold_runs:
        print(f"  P[{i}] style={p.style.name!r}")
        for text, b in bold_runs:
            print(f"    BOLD: {repr(text[:80])}")
        for text, b in normal_runs:
            print(f"    norm: {repr(text[:80])}")

print()
print("=== BOLD TEXT IN TABLES ===")
for ti, tbl in enumerate(doc.tables):
    for ri, row in enumerate(tbl.rows):
        for ci, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                bold_runs = [r.text for r in p.runs if r.font.bold]
                if bold_runs:
                    print(f"  T[{ti}] row={ri} col={ci}: BOLD={[repr(t[:40]) for t in bold_runs]}")
